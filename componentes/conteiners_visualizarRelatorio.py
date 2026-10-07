from pydoc import doc
from xml.dom.minidom import Document

from config import BASE_DIR
from constantes import MODELO_JUVENTUDE, MODELO_PCM, MODELO_SAO_VICENTE
import customtkinter as ctk
import docx
from PIL import Image

from database import buscar_id_usuario, buscar_id_disciplina, listar_modelos
from funcoes_docx.funcoes_leitura import buscarDatas


# ==============================================================================
# CABEÇALHO DA TELA
# ==============================================================================

def criarHeader(quadro_central):
    quadroHeader = ctk.CTkFrame(quadro_central, width=1500, height=150, fg_color="transparent")
    quadroHeader.pack(fill="x", pady=(10, 0), padx=2)

    criar_hearderImg(quadroHeader)
    criar_frameTextos(quadroHeader)

    return quadroHeader


def criar_hearderImg(quadroHeaderl):
    img = ctk.CTkImage(
        light_image=Image.open("imagens/iconeVerRelatorio.png"),
        dark_image=Image.open("imagens/iconeVerRelatorio.png"),
        size=(100, 100)
    )
    labelImg = ctk.CTkLabel(quadroHeaderl, text="", image=img)
    labelImg.pack(side="left", padx=(10, 0), pady=(10, 0))


def criarLabelTitulo(quadroHeader):
    ctk.CTkLabel(
        quadroHeader,
        text="Visualizar Relatórios",
        font=("Roboto", 45),
        text_color="#F8FAFC"
    ).pack(anchor="w")


def criar_TxtHeader(quaadroHeader):
    ctk.CTkLabel(
        quaadroHeader,
        text="Visualize os conteúdos cadastrados no relatório.",
        font=("Roboto", 20),
        text_color="#94A3B8"
    ).pack(anchor="w", padx=(45, 0))


def linhaDivisoria(quadro_principal):
    linhaImg = ctk.CTkImage(
        light_image=Image.open("imagens/linhaCentralGrande.png"),
        dark_image=Image.open("imagens/linhaCentralGrande.png"),
        size=(1450, 80)
    )
    ctk.CTkLabel(quadro_principal, text="", image=linhaImg).pack(fill="x", padx=20)


def criar_frameTextos(quadroHeader):
    frameTexto = ctk.CTkFrame(quadroHeader, fg_color="transparent")
    frameTexto.pack(side="left")

    criarLabelTitulo(frameTexto)
    criar_TxtHeader(frameTexto)

    return frameTexto


# ==============================================================================
# ÁREA DOS FILTROS (MODELO E DATA)
# ==============================================================================

def criar_filtros(quadro_central, usuario, organizaTabela, frameTabela, disciplina):
    frameFiltro = ctk.CTkFrame(quadro_central, fg_color="transparent")
    frameFiltro.pack(fill="x", pady=(20, 0), padx=20)

    frameFiltro.grid_columnconfigure(0, weight=1)
    frameFiltro.grid_columnconfigure(1, weight=1)

    labelModeloRelatorio(frameFiltro)
    labelData(frameFiltro)

    menuModelo = criar_optionModelo(frameFiltro, usuario, disciplina)
    menuData = criar_optionData(frameFiltro, usuario, menuModelo, organizaTabela, frameTabela, disciplina)

    menuModelo.configure(
        command=lambda escolha: atualizaOptionData(
            menuData,
            usuario,
            escolha,
            disciplina
        )
    )


def labelModeloRelatorio(frameModelo):
    modeloRelatorio = ctk.CTkLabel(
        frameModelo,
        text="Modelo do Relatório",
        font=("Roboto", 18, "bold"),
        text_color="#F8FAFC"
    )
    modeloRelatorio.grid(row=0, column=0, sticky="nw", pady=(0, 8), padx=(20, 0))
    return modeloRelatorio


def labelData(frame):
    lbl_data = ctk.CTkLabel(
        frame,
        text="Data do Relatório",
        font=("Roboto", 18, "bold"),
        text_color="#F8FAFC"
    )
    lbl_data.grid(row=0, column=1, pady=(0, 8), padx=(20, 0), sticky="nw")


def criar_btnEnviarDrive(quadro_central):
    """Cria o botão de envio no rodapé da tela, seguindo a lógica da tela de cadastro."""
    btn_drive = ctk.CTkButton(
        quadro_central,
        text="Enviar para o Drive",
        font=("Roboto", 18, "bold"),
        height=52,
        corner_radius=10,
        fg_color="#2563EB",
        hover_color="#1D4ED8",
        text_color="#FFFFFF"
        # Sem comando atribuído: a lógica será implementada por você!
    )
    btn_drive.pack(side="bottom", fill="x", padx=40, pady=(15, 20))
    return btn_drive


def criar_optionModelo(frame, usuario, disciplina):
    id_usuario = buscar_id_usuario(usuario)
    id_disciplina = buscar_id_disciplina(id_usuario, disciplina)

    modelos = listar_modelos(id_disciplina)
    menuModelo = ctk.CTkOptionMenu(
        frame,
        fg_color="#0F172A",
        button_color="#1E293B",
        button_hover_color="#334155",
        text_color="#F8FAFC",
        corner_radius=8,
        height=42,
        values=modelos,
        font=("Roboto", 15),
        dropdown_fg_color="#0F172A",
        dropdown_hover_color="#1E293B",
        dropdown_font=("Roboto", 14)
    )
    menuModelo.grid(row=1, column=0, sticky="w", pady=(0, 10), padx=(20, 0))
    menuModelo.set("Selecione o modelo")
    return menuModelo


def criar_optionData(frame, usuario, menuModelo, organizaTabela, frameTabela, disciplina):
    menuData = ctk.CTkOptionMenu(
        frame,
        fg_color="#0F172A",
        button_color="#1E293B",
        button_hover_color="#334155",
        text_color="#F8FAFC",
        corner_radius=8,
        height=42,
        values=[],
        command=lambda escolha: pegaDataSelecionada(
            escolha, usuario, menuModelo, organizaTabela, frameTabela, disciplina
        ),
        font=("Roboto", 15),
        dropdown_fg_color="#0F172A",
        dropdown_hover_color="#1E293B",
        dropdown_font=("Roboto", 14)
    )
    menuData.grid(row=1, column=1, sticky="w", pady=(0, 10), padx=(20, 0))
    menuData.set("Selecione a data")
    return menuData


def atualizaOptionData(menuData, usuario, modelo, disciplina):
    novas_datas = buscarDatas(usuario, modelo, disciplina)

    menuData.configure(values=novas_datas["datasFormatadas"])

    if novas_datas:
        menuData.set("Selecione a data")


def pegaDataSelecionada(dataSelecionada, usuario, menuModelo, organizaTabela, frameTabela, disciplina):
    modelo = menuModelo.get()

    organizaTabela.pack(
        fill="both",
        expand=True,
        pady=(15, 10),
        padx=20
    )

    frameTabela.pack(fill="both", expand=True)

    limpaTabela(frameTabela)

    cabecalhoTabela(frameTabela, modelo)

    corpoTabela(frameTabela, usuario, modelo, dataSelecionada, disciplina)


# ==============================================================================
# ÁREA DA TABELA DE CONTEÚDO
# ==============================================================================

def tabela(frame):
    organizaTabela = ctk.CTkFrame(frame, fg_color="transparent")

    frameTabela = ctk.CTkScrollableFrame(
        organizaTabela,
        fg_color="#1E293B",
        corner_radius=12,
        border_width=1,
        border_color="#334155"
    )

    frameTabela.grid_columnconfigure(0, weight=1, minsize=420)
    frameTabela.grid_columnconfigure(1, weight=3, minsize=750)

    return organizaTabela, frameTabela


def cabecalhoTabela(frame, modelo):
    texto_coluna1 = "DISCIPLINA / TURMA" if modelo == MODELO_PCM else "DISCIPLINA"

    turmaDisciplina = ctk.CTkLabel(
        frame,
        text=texto_coluna1,
        fg_color="#0F172A",
        text_color="#38BDF8",
        border_width=1,
        border_color="#334155",
        corner_radius=6,
        font=("Roboto", 15, "bold"),
        height=48
    )
    turmaDisciplina.grid(row=0, column=0, sticky="ew", padx=8, pady=(12, 6))

    conteudo = ctk.CTkLabel(
        frame,
        text="CONTEÚDO MINISTRADO",
        fg_color="#0F172A",
        text_color="#38BDF8",
        border_width=1,
        border_color="#334155",
        corner_radius=6,
        font=("Roboto", 15, "bold"),
        height=48
    )
    conteudo.grid(row=0, column=1, sticky="ew", padx=8, pady=(12, 6))


def corpoTabela(frame, usuario, modelo, data, disciplina):
    dataFormatada = data.replace("/", "-")
    caminho_documento = (
        BASE_DIR / "Professores" / usuario / disciplina / f"backup{modelo}" /
        f"Relatorio{modelo}{disciplina}{dataFormatada}.docx"
    )

    if not caminho_documento.exists():
        return

    doc = docx.Document(caminho_documento)
    cont = 1

    if modelo == MODELO_PCM:
        for tabela_doc in doc.tables:
            for linha in tabela_doc.rows:
                if "TURMA" in linha.cells[0].text:
                    turma = ctk.CTkLabel(
                        frame,
                        text=linha.cells[0].text.strip(),
                        fg_color="#16203B",
                        text_color="#F8FAFC",
                        border_width=1,
                        border_color="#334155",
                        corner_radius=6,
                        font=("Roboto", 16, "bold"),
                        height=64
                    )
                    turma.grid(row=cont, column=0, sticky="nsew", padx=8, pady=4)

                    conteudo = ctk.CTkLabel(
                        frame,
                        text=linha.cells[1].text.strip(),
                        fg_color="#0F172A",
                        text_color="#CBD5E1",
                        border_width=1,
                        border_color="#334155",
                        corner_radius=6,
                        font=("Roboto", 13),
                        height=64,
                        anchor="w",
                        padx=18,
                        wraplength=720,
                        justify="left"
                    )
                    conteudo.grid(row=cont, column=1, sticky="nsew", padx=8, pady=4)

                    cont += 1

    elif modelo == MODELO_JUVENTUDE:
        for tabela_doc in doc.tables:
            for linha in tabela_doc.rows:
                if "PRÁTICA DE CONJUNTO" in linha.cells[0].text:
                    turma = ctk.CTkLabel(
                        frame,
                        text=linha.cells[0].text.strip(),
                        fg_color="#16203B",
                        text_color="#F8FAFC",
                        border_width=1,
                        border_color="#334155",
                        corner_radius=6,
                        font=("Roboto", 16, "bold"),
                        height=64
                    )
                    turma.grid(row=cont, column=0, sticky="nsew", padx=8, pady=4)

                    conteudo = ctk.CTkLabel(
                        frame,
                        text=linha.cells[1].text.strip(),
                        fg_color="#0F172A",
                        text_color="#CBD5E1",
                        border_width=1,
                        border_color="#334155",
                        corner_radius=6,
                        font=("Roboto", 13),
                        height=64,
                        anchor="w",
                        padx=18,
                        wraplength=720,
                        justify="left"
                    )
                    conteudo.grid(row=cont, column=1, sticky="nsew", padx=8, pady=4)

                    cont += 1

    elif modelo == MODELO_SAO_VICENTE:
        pass


def limpaTabela(frameTabela):
    for widget in frameTabela.winfo_children():
        widget.destroy()
